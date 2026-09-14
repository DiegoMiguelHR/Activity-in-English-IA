# -*- coding: utf-8 -*-
"""================================================================================
    BUILD REPORT (ENGLISH): creates the Word report for the project
    "Activity in English" - Comparison of Classification Models.

    Content:
      - Cover page (faculty, topic, course, teacher, members, city-year)
      - Phase 1: Theoretical Matrix
      - Phase 2: The Classification Challenge (dataset, pipeline, optimization,
        cross-validation, results and figures)
      - Phase 3: The Verdict (final comparative analysis)
      - Annex: Reproducibility

    USAGE:
        python generar_informe_en.py
================================================================================
"""

import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.abspath(__file__))
GRAF = os.path.join(BASE, "ActividadIngles", "graficos")
OUT_PATH = os.path.join(BASE, "Activity in English - FINAL REPORT.docx")

AZUL = RGBColor(0x1F, 0x4E, 0x79)


def configurar_documento(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    for level, size in [("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 12)]:
        st = doc.styles[level]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = AZUL
        st.font.bold = True
        st.element.get_or_add_rPr()


def parrafo_justificado(doc, texto):
    p = doc.add_paragraph(texto)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def add_tabla(doc, encabezados, filas, anchos=None):
    tabla = doc.add_table(rows=1 + len(filas), cols=len(encabezados))
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER

    for j, h in enumerate(encabezados):
        celda = tabla.cell(0, j)
        celda.text = ""
        run = celda.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(10)
        shd = celda._tc.get_or_add_tcPr()
        from docx.oxml import OxmlElement
        el = OxmlElement("w:shd")
        el.set(qn("w:val"), "clear")
        el.set(qn("w:color"), "auto")
        el.set(qn("w:fill"), "DCE6F1")
        shd.append(el)

    for i, fila in enumerate(filas):
        for j, valor in enumerate(fila):
            celda = tabla.cell(i + 1, j)
            celda.text = ""
            run = celda.paragraphs[0].add_run(str(valor))
            run.font.size = Pt(10)

    if anchos:
        for j, ancho in enumerate(anchos):
            for i in range(len(filas) + 1):
                tabla.cell(i, j).width = Inches(ancho)
    return tabla


def add_figura(doc, ruta, titulo, ancho=6.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(ruta, width=Inches(ancho))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(titulo)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)


def add_figura_en_celda(celda, ruta, titulo, ancho=2.9):
    p = celda.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(ruta, width=Inches(ancho))
    cap = celda.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(titulo)
    r.italic = True
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)


def add_codigo(doc, codigo):
    for linea in codigo.strip("\n").splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(linea if linea else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")


doc = Document()
configurar_documento(doc)

# ============================ COVER PAGE =====================================
for _ in range(4):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("UNIVERSIDAD ANDINA DEL CUSCO")
r.bold = True
r.font.size = Pt(22)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("FACULTAD DE INGENIERÍA Y ARQUITECTURA")
r.bold = True
r.font.size = Pt(15)

for _ in range(5):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('"Activity in English"')
r.bold = True
r.font.size = Pt(26)
r.font.color.rgb = AZUL

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Comparison of Classification Models with Supervised Learning")
r.font.size = Pt(13)

for _ in range(4):
    doc.add_paragraph()

datos = [
    ("COURSE:", "Artificial Intelligence"),
    ("TEACHER:", "Espetia Huamanga, Hugo"),
]
for etiqueta, valor in datos:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(etiqueta + " ")
    r.bold = True
    r.font.size = Pt(12)
    r = p.add_run(valor)
    r.font.size = Pt(12)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("MEMBERS:")
r.bold = True
r.font.size = Pt(12)

integrantes = [
    "Huamantalla Ríos, Diego Miguel   -   022100354D",
    "Latorre Campos, Andre Ismael     -   022101034C",
    "Ordóñez Paniura, Marcela Leonor  -   022100735H",
    "Suel Monrroy, David Joshua       -   022101091G",
]
for nombre in integrantes:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(nombre)
    r.font.size = Pt(12)

for _ in range(4):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("CUSCO – PERU")
r.bold = True
r.font.size = Pt(13)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("2026")
r.bold = True
r.font.size = Pt(13)

doc.add_page_break()

# ===================== PROJECT PRESENTATION ==================================
doc.add_heading("Project Presentation", level=1)

parrafo_justificado(
    doc,
    "This work applies four supervised classification models to a real-world problem based on "
    "environmental sensor data. The activity is organized in three phases: (1) strategic "
    "selection and theoretical analysis of the models, (2) classification challenge with "
    "experimentation and rigorous validation, and (3) final comparative verdict on the best "
    "model in terms of performance and computational cost.",
)
parrafo_justificado(
    doc,
    "The four selected models belong to different mathematical families, as required by the "
    "activity: Support Vector Machines (SVM) and K-Nearest Neighbors (KNN) from the geometric "
    "and distance-based family; Gaussian Naive Bayes from the probabilistic family; and Gradient "
    "Boosting from the ensemble family. The chosen dataset was the \u201cOccupancy Detection "
    "Dataset\u201d, on which all four algorithms competed on equal terms through a unified "
    "preprocessing pipeline.",
)

# ============================ PHASE 1 ========================================
doc.add_heading("Phase 1: Theoretical Matrix", level=1)

parrafo_justificado(
    doc,
    "According to the activity rules, each team had to select four models ensuring at least one "
    "model from each mathematical family. The selected combination was:",
)
parrafo_justificado(
    doc,
    "\u2022  Family 1 (Geometric and distance-based): Support Vector Machines (SVM) and "
    "K-Nearest Neighbors (KNN), evaluating different kernels and distance metrics.\n"
    "\u2022  Family 2 (Probabilistic and generative): Gaussian Naive Bayes (GaussianNB).\n"
    "\u2022  Family 3 (Ensemble and stochastic): Gradient Boosting.",
)
parrafo_justificado(
    doc,
    "Table 1 contrasts the four models according to the four requested criteria: core "
    "mathematical principle, sensitivity to outliers, computational cost, and assumptions about "
    "data distribution.",
)

add_tabla(
    doc,
    ["Model", "Core Mathematical Principle", "Sensitivity to Outliers",
     "Computational Cost", "Assumptions about Data Distribution"],
    [
        ["Support Vector Machines (SVM)",
         "Margin maximization (distance) between classes using hyperplanes. Applies the \u201ckernel trick\u201d to project the data into higher dimensions.",
         "Low to Medium. Strongly depends on the support vectors: if an outlier lies close to the margin, it can shift the hyperplane.",
         "High. Training scales quadratically with the number of instances (O(n\u00b2)), and requires tuning C, gamma and the kernel.",
         "Assumes no specific underlying distribution (non-parametric model); it only depends on the geometry of the data."],
        ["Gaussian Naive Bayes (GNB)",
         "Bayes\u2019 Theorem combined with the assumption of conditional independence among the predictor variables.",
         "Moderate to High. When estimating means and variances, extreme values can bias the parameters of the distributions.",
         "Very Low. Highly efficient in both training and prediction; requires few resources.",
         "Assumes that the numeric features follow a Normal (Gaussian) distribution and that all variables are independent of each other."],
        ["Gradient Boosting",
         "Stochastic and sequential ensemble: it iteratively combines weak decision trees, where each new tree minimizes the residual errors of the previous one through gradient descent.",
         "Low. It is inherently robust to outliers thanks to the base trees and robust loss functions.",
         "Very High. Trains multiple trees sequentially (not natively parallelizable like Random Forest), demanding much time and memory.",
         "Makes no mathematical assumptions about the data distribution or the spatial relationship among variables."],
        ["K-Nearest Neighbors (KNN)",
         "Classification based on spatial distance (Euclidean, Manhattan, etc.) to the k nearest neighbors in feature space.",
         "High. Being distance-based, a single outlier can erroneously become the \u201cnearest neighbor\u201d of a new sample.",
         "Low in training, high in prediction (lazy learning). Computing distances against all instances is costly in large datasets.",
         "Assumes no prior distribution, but does assume that samples with similar features belong to the same class (spatial locality)."],
    ],
    anchos=[1.05, 1.75, 1.45, 1.50, 1.75],
)

cap = doc.add_paragraph("Table 1. Comparative theoretical matrix of the four classification models.")
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap.runs[0].italic = True
cap.runs[0].font.size = Pt(9)

doc.add_page_break()

# ============================ PHASE 2 ========================================
doc.add_heading("Phase 2: The Classification Challenge", level=1)

# ---- 2.1 Dataset ----
doc.add_heading("2.1. Dataset: Occupancy Detection", level=2)

parrafo_justificado(
    doc,
    "The \u201cOccupancy Detection Dataset\u201d was selected, an environmental sensor data "
    "classification problem (the suggested category of \u201cenvironmental sensor data\u201d). "
    "The goal is to predict whether a room is occupied (1) or empty (0) from sensor readings "
    "recorded approximately every minute.",
)
parrafo_justificado(
    doc,
    "The original dataset is distributed in two files (training and test). For this experiment "
    "both were concatenated, obtaining a total of 17,895 instances, well above the required "
    "minimum of 2,500. Two non-predictive columns were removed: the unnamed index column "
    "(\u201cUnnamed: 0\u201d) and the date/time column (\u201cdate\u201d), since their use could "
    "introduce temporal artifacts and information leakage.",
)

add_tabla(
    doc,
    ["Variable", "Description", "Type"],
    [
        ["Temperature", "Ambient temperature (\u00b0C)", "Numeric predictor"],
        ["Humidity", "Relative humidity (%)", "Numeric predictor"],
        ["Light", "Light intensity (lux)", "Numeric predictor"],
        ["CO2", "Carbon dioxide in the air (ppm)", "Numeric predictor"],
        ["HumidityRatio", "Humidity ratio (kg water / kg dry air)", "Numeric predictor"],
        ["Occupancy", "Room state: 0 = empty, 1 = occupied", "Binary target"],
    ],
    anchos=[1.6, 3.4, 1.7],
)
cap = doc.add_paragraph("Table 2. Dataset variables. Five numeric predictors and one binary target variable.")
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap.runs[0].italic = True
cap.runs[0].font.size = Pt(9)

doc.add_paragraph()
p = doc.add_paragraph()
r = p.add_run("Class distribution: ")
r.bold = True
parrafo_justificado(
    doc,
    "14,117 instances of empty room (78.9%) versus 3,778 of occupied room (21.1%). The dataset is "
    "imbalanced, so accuracy would not be a reliable metric: a classifier that always predicted "
    "\u201cempty\u201d would reach 78.9% accuracy without learning anything. Therefore, the main "
    "evaluation metric is the average (macro) F1-Score.",
)

# ---- 2.2 Pipeline ----
doc.add_heading("2.2. Unified preprocessing pipeline", level=2)
parrafo_justificado(
    doc,
    "So that the four models compete on equal terms, a single pipeline was built consisting of a "
    "StandardScaler (z-score normalization) followed by the classifier. Scaling is essential for "
    "distance-based models (KNN) and margin-based models (SVM), since variables such as light "
    "(hundreds of lux) and humidity ratio (units of 10\u207b\u00b3) are on very different scales.",
)
parrafo_justificado(
    doc,
    "Importantly, the StandardScaler sits inside the Pipeline passed to GridSearchCV: in this "
    "way, the scaling statistics are learned only with the training folds of each validation, "
    "avoiding data leakage.",
)
add_codigo(doc, """
    pipeline = Pipeline([
        ("scaler", StandardScaler()),
        ("clf", classifier),
    ])
""")

# ---- 2.3 Optimization ----
doc.add_heading("2.3. Training and hyperparameter optimization", level=2)
parrafo_justificado(
    doc,
    "Each model was optimized with GridSearchCV, an exhaustive search over the hyperparameter "
    "grids shown in Table 3, using the macro F1-Score as the search metric. For SVM, three "
    "different kernels were evaluated (kernel analysis) and for KNN two distance metrics "
    "(distance-metric evaluation), as required by the geometric family.",
)

add_tabla(
    doc,
    ["Model", "Hyperparameter grid (GridSearch)", "# combinations"],
    [
        ["SVC", "kernel: [rbf, linear, poly] \u00b7 C: [0.1, 1, 10] \u00b7 gamma: [0.01, 0.1, 1]", "27"],
        ["GaussianNB", "var_smoothing: [1e-9, 1e-8, 1e-7]", "3"],
        ["GradientBoosting", "n_estimators: [50, 100] \u00b7 max_depth: [2, 4] \u00b7 learning_rate: [0.05, 0.1]", "8"],
        ["KNN", "n_neighbors: [3, 5, 7] \u00b7 weights: [uniform, distance] \u00b7 metric: [euclidean, manhattan]", "12"],
    ],
    anchos=[1.5, 3.6, 1.6],
)
cap = doc.add_paragraph("Table 3. Hyperparameter grids evaluated by GridSearchCV for each model.")
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap.runs[0].italic = True
cap.runs[0].font.size = Pt(9)

# ---- 2.4 Validation ----
doc.add_heading("2.4. Cross-validation", level=2)
parrafo_justificado(
    doc,
    "To guarantee that results do not depend on a lucky data split, stratified cross-validation "
    "with 5 folds was used (StratifiedKFold, 5 splits, shuffle=True, random_state=42). The "
    "stratified version keeps the original class proportion (78.9% / 21.1%) in every fold, which "
    "is the most rigorous choice given the class imbalance. The reported F1-Score is the average "
    "over the five folds.",
)

# ---- 2.5 Results ----
doc.add_heading("2.5. Experimental results", level=2)
parrafo_justificado(
    doc,
    "Table 4 summarizes the final performance of each model after optimization. KNN achieved the "
    "best macro F1-Score (0.9923) with the lowest computational load among the top three. The "
    "best hyperparameters found were:",
)

add_tabla(
    doc,
    ["Model", "F1-Score (macro)", "Accuracy", "Precision (macro)", "Recall (macro)",
     "Total time (s)", "Best hyperparameters"],
    [
        ["KNN", "0.9923", "0.9949", "0.9912", "0.9934", "1.21",
         "n_neighbors=5, weights='distance', metric='euclidean'"],
        ["SVC", "0.9900", "0.9933", "0.9869", "0.9931", "34.51",
         "C=10, gamma=1, kernel='rbf'"],
        ["GradientBoosting", "0.9886", "0.9923", "0.9866", "0.9906", "8.75",
         "n_estimators=100, max_depth=4, learning_rate=0.1"],
        ["GaussianNB", "0.9530", "0.9670", "0.9327", "0.9784", "0.17",
         "var_smoothing=1e-9"],
    ],
    anchos=[1.4, 0.85, 0.7, 0.95, 0.9, 0.85, 1.9],
)
cap = doc.add_paragraph("Table 4. Final results: average macro F1-Score from cross-validation (StratifiedKFold, 5 folds). Precision/recall/accuracy were estimated with out-of-fold predictions.")
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap.runs[0].italic = True
cap.runs[0].font.size = Pt(9)

doc.add_paragraph()
add_figura(doc, os.path.join(GRAF, "resultados_f1.png"),
           "Figure 1. Average macro F1-Score from cross-validation per model.", 5.6)
add_figura(doc, os.path.join(GRAF, "resultados_tiempo.png"),
           "Figure 2. Total computational time (search + training) per model.", 5.6)

doc.add_page_break()

doc.add_heading("2.6. Kernel analysis in SVM and distance metrics in KNN", level=2)

parrafo_justificado(
    doc,
    "As required by the geometric family, different kernels were analyzed for SVM. The RBF kernel "
    "was clearly superior, followed by the linear and then the polynomial kernel. For KNN, the "
    "Euclidean metric barely outperformed Manhattan, indicating that the geometric structure of "
    "the problem favors straight-line distance.",
)
add_figura(doc, os.path.join(GRAF, "comparacion_kernels.png"),
           "Figure 3. Best F1-Score achieved by each SVM kernel.", 5.4)
add_figura(doc, os.path.join(GRAF, "comparacion_metricas.png"),
           "Figure 4. Best F1-Score achieved by each KNN distance metric.", 5.4)

doc.add_heading("2.7. Confusion matrices", level=2)
parrafo_justificado(
    doc,
    "The confusion matrices were built with out-of-fold predictions (each instance was predicted "
    "by a model trained without it). They make it possible to quantify the errors of each model:",
)

tabla_cm = doc.add_table(rows=2, cols=2)
tabla_cm.style = "Table Grid"
tabla_cm.alignment = WD_TABLE_ALIGNMENT.CENTER
add_figura_en_celda(tabla_cm.cell(0, 0), os.path.join(GRAF, "confusion_KNN.png"),
                    "KNN \u2014 errors: 58 FP, 34 FN")
add_figura_en_celda(tabla_cm.cell(0, 1), os.path.join(GRAF, "confusion_SVC.png"),
                    "SVC \u2014 errors: 93 FP, 27 FN")
add_figura_en_celda(tabla_cm.cell(1, 0), os.path.join(GRAF, "confusion_GradientBoosting.png"),
                    "GradientBoosting \u2014 errors: 90 FP, 47 FN")
add_figura_en_celda(tabla_cm.cell(1, 1), os.path.join(GRAF, "confusion_GaussianNB.png"),
                    "GaussianNB \u2014 errors: 584 FP, 7 FN")
doc.add_paragraph()
cap = doc.add_paragraph("Figures 5\u20138. Confusion matrices of the four models (out-of-fold predictions). FP = false positives, FN = false negatives.")
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap.runs[0].italic = True
cap.runs[0].font.size = Pt(9)

doc.add_heading("2.8. Interactive application (Streamlit)", level=2)
parrafo_justificado(
    doc,
    "Additionally, a web interface was developed with Streamlit (app_occupancy.py) that "
    "replicates the experiment and runs it at the click of a button, displaying the results "
    "table, the best hyperparameters per model and an interactive F1-Score bar chart. It "
    "facilitated capturing orderly screenshots for the report.",
)
add_figura(doc, os.path.join(GRAF, "captura_streamlit.png"),
           "Figure 9. Screenshot of the Streamlit app showing the model comparison.", 5.0)

doc.add_page_break()

# ============================ PHASE 3 ========================================
doc.add_heading("Phase 3: The Verdict", level=1)
parrafo_justificado(
    doc,
    "Based on the experimental results, the three final analysis questions are answered below.",
)

doc.add_heading("3.1. Balance between accuracy and generalization capability (F1-Score)", level=2)
parrafo_justificado(
    doc,
    "After stratified cross-validation, the model that showed the best balance between accuracy "
    "and generalization was KNN, with an outstanding macro F1-Score of 0.9923 and an accuracy of "
    "0.9949. This metric balances almost perfectly false positives and false negatives when "
    "predicting room occupancy (58 FP and 34 FN out of 17,895 instances), beating by a "
    "noticeable margin the lowest-performing model, Gaussian Naive Bayes, which reached 0.9530 "
    "(with 584 false positives: it predicted \u201coccupied\u201d on 584 empty-room readings). "
    "Using the macro F1 guarantees that the result is not distorted by the class imbalance "
    "(78.9% of empty rooms).",
)

doc.add_heading("3.2. Did the most complex model justify its computational cost?", level=2)
parrafo_justificado(
    doc,
    "No. In this experiment, the mathematically more complex models did NOT justify their high "
    "computational cost at all. While SVC (34.51 s) and Gradient Boosting (8.75 s) reached "
    "slightly lower F1-Scores (0.9900 and 0.9886), KNN obtained the winning score (0.9923) by "
    "running the whole process in just 1.21 seconds, and GaussianNB did it in 0.17 s. The "
    "simplicity of computing Euclidean distances clearly beat the effort of building sequential "
    "trees or projecting multidimensional hyperplanes. The extra cost of SVC and Boosting did not "
    "translate into better performance, so for this problem the additional complexity is not "
    "profitable.",
)
add_figura(doc, os.path.join(GRAF, "resultados_f1_tiempo.png"),
           "Figure 10. Performance (F1-Score) versus computational cost: the optimal quadrant is top-left.", 5.6)

doc.add_heading("3.3. Decision boundaries and problem dimensionality", level=2)
parrafo_justificado(
    doc,
    "The dataset has five environmental variables (Temperature, Humidity, Light, CO2, "
    "HumidityRatio). Projecting the data onto its two first PCA components retains about 82% of "
    "the variance (46.4% + 35.8%), which indicates an intrinsically low-dimensional structure. "
    "In that projection (Figure 11), the classes (empty vs. occupied room) form dense and "
    "spatially separated clusters: when people are present, light and CO2 rise simultaneously, "
    "grouping the occupied-class points.",
)
parrafo_justificado(
    doc,
    "This structure explains KNN\u2019s triumph: it drew decision boundaries based purely on data "
    "proximity, without rigid assumptions, and its frontier almost coincides with the natural "
    "separation of the clusters. SVC (RBF kernel) and Gradient Boosting achieved almost "
    "equivalent boundaries, but at a much higher computational cost. On the opposite end, "
    "GaussianNB drew an overly rigid probabilistic boundary: by assuming independence and "
    "normality among variables that are actually correlated (Light and CO2, because of their "
    "link to human presence; Temperature and HumidityRatio), its boundary was displaced and "
    "produced the highest number of false positives of the four models. The low effective "
    "dimensionality prevented the curse of dimensionality, favoring distance-based techniques "
    "over the parametric frontiers of Naive Bayes.",
)
add_figura(doc, os.path.join(GRAF, "fronteras_decision_PCA.png"),
           "Figure 11. Decision boundaries of the four models projected onto the two first PCA components.", 6.2)
add_figura(doc, os.path.join(GRAF, "correlacion_features.png"),
           "Figure 12. Correlation among the predictor variables (supports the analysis on Gaussian Naive Bayes).", 4.8)

doc.add_heading("3.4. Final conclusion", level=2)
parrafo_justificado(
    doc,
    "KNN with k=5, distance weighting and the Euclidean metric was the winning model, achieving "
    "the best F1-Score with the lowest computational cost among the three high-performing models. "
    "The complex models (SVC and Gradient Boosting) did not justify their extra cost in this "
    "problem, and the probabilistic model (GaussianNB) lagged behind because its independence "
    "assumptions were violated. The practical lesson: for low-dimensional problems with "
    "well-separated classes, a simple geometric method such as KNN can be the most efficient and "
    "accurate option.",
)

doc.add_page_break()

# ============================ ANNEX ==========================================
doc.add_heading("Annex: Reproducibility", level=1)

parrafo_justificado(
    doc,
    "All experiment files are located in the ActividadIngles folder. The results presented in "
    "this report were generated with the following procedure:",
)

doc.add_heading("A.1. Code and execution", level=2)
add_tabla(
    doc,
    ["File", "Description", "Command"],
    [
        ["occupancy_classification.py", "Complete console experiment (data, pipeline, GridSearchCV, results)", "python occupancy_classification.py"],
        ["app_occupancy.py", "Interactive Streamlit app of the experiment", "streamlit run app_occupancy.py"],
        ["generar_graficos.py", "Generates all the figures of this report in ./graficos", "python generar_graficos.py"],
        ["generar_informe_en.py", "Generates this final Word document", "python generar_informe_en.py"],
    ],
    anchos=[2.1, 3.4, 1.9],
)
cap = doc.add_paragraph("Table 5. Project files and commands to reproduce each part.")
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap.runs[0].italic = True
cap.runs[0].font.size = Pt(9)

doc.add_heading("A.2. Dependencies", level=2)
add_codigo(doc, """
    numpy, pandas, scikit-learn, matplotlib, streamlit, python-docx
""")

doc.add_heading("A.3. Experimentation flow summary", level=2)
add_codigo(doc, """
    1. Load and concatenate DataTraining.csv + DataTest.csv  -> 17,895 rows
    2. Cleanup: remove 'Unnamed: 0' and 'date'
    3. Unified pipeline: StandardScaler + classifier
    4. GridSearchCV (scoring = macro F1, cv = StratifiedKFold(5))
    5. Sort results by F1-Score and computational time
""")

doc.save(OUT_PATH)
print("Report generated:", OUT_PATH)